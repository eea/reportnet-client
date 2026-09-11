"""Read a public dataflow's schema and quality-control rules — no API key needed.

When a dataflow is published, Reportnet serves its whole design as a spreadsheet
at ``GET /dataflow/downloadPublicSchemaInformation/{dataflowId}``, unauthenticated.
That file carries every QC rule, **including the SQL**, which the ``/rules/``
endpoints refuse to an API key of any role.

So this is the only way to read a dataflow's quality rules programmatically:

    >>> schema = reportnet.get_public_schema(1748)
    >>> len(schema.rules)
    843
    >>> sql = [r for r in schema.rules if r.is_sql]
    >>> sql[0].shortcode, sql[0].severity
    ('ReportPeriod-4', 'WARNING')

Rules address data as ``dataset_<id>."table"`` where ``<id>`` is a *design*
dataset. :attr:`QcRule.dataset_ids` pulls those ids out; resolve them with
:meth:`~reportnet.DataflowClient.get_design_datasets` (which does need a key).

Only public dataflows are served — anything else returns 404, including private
test copies of a published dataflow.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

import httpx

from ._log import get_logger

logger = get_logger(__name__)

PUBLIC_SCHEMA_PATH = "/dataflow/downloadPublicSchemaInformation"

_DATASET_REF = re.compile(r'dataset_(\d+)\s*\.\s*"?(\w+)"?', re.IGNORECASE)
_SQL_START = re.compile(r"^\s*(select|with)\b", re.IGNORECASE)
_HEADER = ("Table", "Field", "Shortcode")


@dataclass(frozen=True)
class QcRule:
    """One quality-control rule, as published in a dataflow's schema export."""

    dataset: str          # which dataset in the dataflow, e.g. "Descriptive data"
    table: str
    field: str            # "" for table- and record-level rules
    shortcode: str
    name: str
    description: str
    expression: str       # SQL, or Reportnet's own expression language, or ""
    qc_type: str          # FIELD | RECORD | TABLE
    severity: str         # BLOCKER | ERROR | WARNING | INFO
    message: str
    automatic: bool
    enabled: bool

    @property
    def is_sql(self) -> bool:
        """True when the expression is SQL rather than the expression language.

        Automatic rules (cardinality, uniqueness, type checks) carry no
        expression at all; the rest are either SQL or Reportnet's own
        ``( field MATCH ^regex$ )`` syntax.
        """
        return bool(_SQL_START.match(self.expression))

    @property
    def dataset_ids(self) -> tuple[int, ...]:
        """Design dataset ids this rule's SQL reads from, in order of appearance."""
        seen: dict[int, None] = {}
        for raw, _table in _DATASET_REF.findall(self.expression):
            seen.setdefault(int(raw), None)
        return tuple(seen)

    @property
    def tables_read(self) -> tuple[str, ...]:
        """Table names this rule's SQL reads, lowercased as Reportnet stores them."""
        seen: dict[str, None] = {}
        for _id, table in _DATASET_REF.findall(self.expression):
            seen.setdefault(table.lower(), None)
        return tuple(seen)


@dataclass(frozen=True)
class PublicSchema:
    """A public dataflow's design: its QC rules, plus the spreadsheet they came from."""

    dataflow_id: int
    content: bytes                  # the .xlsx exactly as served
    filename: str                   # the name Reportnet gave it
    rules: tuple[QcRule, ...]

    def save(self, path: str | Path) -> Path:
        """Write the original spreadsheet to *path* (a file, or a directory)."""
        target = Path(path)
        if target.is_dir():
            target = target / self.filename
        target.write_bytes(self.content)
        return target

    def sql_rules(self) -> tuple[QcRule, ...]:
        """Only the rules with SQL — the ones a custodian wrote by hand."""
        return tuple(r for r in self.rules if r.is_sql)

    def by_dataset(self) -> dict[str, tuple[QcRule, ...]]:
        """Rules grouped by the dataset they belong to."""
        out: dict[str, list[QcRule]] = {}
        for rule in self.rules:
            out.setdefault(rule.dataset, []).append(rule)
        return {k: tuple(v) for k, v in out.items()}


def _yes(value: str) -> bool:
    return value.strip().lower() in ("yes", "true", "1")


def _rows(sheet: Any) -> Iterator[list[str]]:
    for row in sheet.iter_rows(values_only=True):
        yield ["" if cell is None else str(cell).strip() for cell in row]


def _parse_rules(content: bytes) -> tuple[QcRule, ...]:
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError:
        raise ImportError(
            "openpyxl is required to parse a public schema; install it with: "
            'pip install "reportnet-client[schema]"'
        ) from None

    import io
    import warnings

    with warnings.catch_warnings():
        # Reportnet's export has no default style; openpyxl warns and carries on.
        warnings.simplefilter("ignore")
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if "QC rules" not in workbook.sheetnames:
            return ()
        dataset = ""
        header: list[str] | None = None
        rules: list[QcRule] = []
        for cells in _rows(workbook["QC rules"]):
            if not any(cells):
                continue
            # A lone value in the first column names the dataset the next rules belong to.
            if cells[0] and not any(cells[1:]):
                dataset, header = cells[0], None
                continue
            if tuple(cells[:3]) == _HEADER:
                header = cells
                continue
            if header is None:
                continue
            row = dict(zip(header, cells))
            rules.append(
                QcRule(
                    dataset=dataset,
                    table=row.get("Table", ""),
                    field=row.get("Field", ""),
                    shortcode=row.get("Shortcode", ""),
                    name=row.get("Name", ""),
                    description=row.get("Description", ""),
                    expression=row.get("Expression", ""),
                    qc_type=row.get("Type of QC", ""),
                    severity=row.get("Severity Level", ""),
                    message=row.get("Message", ""),
                    automatic=_yes(row.get("Automatic", "")),
                    enabled=_yes(row.get("Enabled", "")),
                )
            )
        return tuple(rules)
    finally:
        workbook.close()


def get_public_schema(
    dataflow_id: int,
    *,
    base_url: str = "https://api.reportnet.europa.eu",
    timeout: float = 60.0,
    parse: bool = True,
) -> PublicSchema:
    """Download a public dataflow's schema and quality-control rules.

    Needs **no API key** — the endpoint is unauthenticated. It is also the only
    programmatic route to rule SQL: every ``/rules/`` endpoint refuses an API
    key, custodian ones included.

    Args:
        dataflow_id: The dataflow to read.
        base_url: API root; override for sandbox.
        timeout: Seconds to wait for the download.
        parse: Set False to skip parsing and keep only :attr:`PublicSchema.content`
            (no ``openpyxl`` needed).

    Raises:
        :class:`~reportnet.DataflowNotPublicError`: The dataflow is not published.
            Private dataflows are not served here at all, including a private test
            copy of a public one.
        ImportError: *parse* is True and ``openpyxl`` is not installed.

    Example::

        schema = reportnet.get_public_schema(1748)
        for rule in schema.sql_rules():
            print(rule.shortcode, rule.severity, rule.dataset_ids)
    """
    from .exceptions import DataflowNotPublicError

    url = f"{base_url.rstrip('/')}{PUBLIC_SCHEMA_PATH}/{dataflow_id}"
    response = httpx.get(url, timeout=timeout, follow_redirects=True)
    if response.status_code == 404:
        raise DataflowNotPublicError(dataflow_id)
    response.raise_for_status()

    disposition = response.headers.get("content-disposition") or ""
    match = re.search(r"filename=([^;]+)", disposition)
    filename = (
        match.group(1).strip().strip('"')
        if match
        else f"dataflow-{dataflow_id}-Schema_Information.xlsx"
    )
    rules = _parse_rules(response.content) if parse else ()
    logger.debug("public schema for dataflow %s: %d rules", dataflow_id, len(rules))
    return PublicSchema(
        dataflow_id=dataflow_id,
        content=response.content,
        filename=filename,
        rules=rules,
    )
